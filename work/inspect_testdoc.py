from docx import Document
import json
import sys

sys.stdout.reconfigure(encoding="utf-8")

src = r"C:\Users\Quang Hung Computer\Downloads\SU26_StockLite_Report_5.0_TestDoc_v1.7.docx"
doc = Document(src)
matches = []
terms = ("integration", "blocked", "905", "843", "812", "849", "regression", "report 5.2")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text and any(t in text.lower() for t in terms):
        matches.append({"type":"paragraph","index":i,"style":p.style.name,"text":text})
for ti, table in enumerate(doc.tables):
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    joined = " | ".join(" | ".join(r) for r in rows)
    if any(t in joined.lower() for t in terms):
        matches.append({"type":"table","index":ti,"rows":rows})
print(json.dumps({"paragraphs":len(doc.paragraphs),"tables":len(doc.tables),"matches":matches},ensure_ascii=False,indent=2))
