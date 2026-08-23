from docx import Document
import json, sys
sys.stdout.reconfigure(encoding="utf-8")
doc=Document(r"C:\Users\Quang Hung Computer\Downloads\SU26_StockLite_Report_5.0_TestDoc_v1.7.docx")
print(json.dumps({"paragraphs":[{"i":i,"style":p.style.name,"text":p.text} for i,p in enumerate(doc.paragraphs)],"table0":[[c.text for c in r.cells] for r in doc.tables[0].rows]},ensure_ascii=False,indent=2))
