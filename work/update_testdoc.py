from copy import deepcopy
from pathlib import Path
from docx import Document

src = Path(r"C:\Users\Quang Hung Computer\Downloads\SU26_StockLite_Report_5.0_TestDoc_v1.7.docx")
out = Path(r"C:\Users\Quang Hung Computer\Documents\Codex\2026-08-23\b-n-c-n-ti-p\outputs\SU26_StockLite_Report_5.0_TestDoc_v1.8.docx")
out.parent.mkdir(parents=True, exist_ok=True)
doc = Document(src)

def set_cell(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(text)
    for extra in cell.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""

def replace_in_cell(cell, old, new):
    full = "\n".join(p.text for p in cell.paragraphs)
    if old not in full:
        return False
    set_cell(cell, full.replace(old, new))
    return True

# Cover metadata.
set_cell(doc.tables[0].cell(1, 1), "v1.8")
set_cell(doc.tables[0].cell(3, 1), "23/08/2026")

# Add a formatted change-log row while preserving v1.7 as historical evidence.
change_log = doc.tables[1]
new_tr = deepcopy(change_log.rows[-1]._tr)
change_log._tbl.append(new_tr)
new_row = change_log.rows[-1]
for idx, text in enumerate([
    "1.8",
    "23/08/2026",
    "Nguyễn Đình Đức",
    "Synchronized Report 5.2 Integration Test results after final status closure: 905 total cases, 843 Pass, 62 Fail, 0 Blocked and 0 Not Run. Updated controller/API, Scheduler, Security, execution-summary and defect references."
]):
    set_cell(new_row.cells[idx], text)

# Security Testing table: current L2 status and evidence wording.
security = doc.tables[12]
set_cell(security.cell(1, 1), "Verify selected L2 authentication, server-side authorization and security controls against BR-45, BR-46, BR-47, BR-51, BR-55 and NFR-S01–S03. The scope is limited to controls with actual API, xUnit, source-review or log evidence.")
set_cell(security.cell(2, 1), "Use automated Postman/HTTP requests against the local ASP.NET Core API for missing or malformed JWTs, wrong-role access, HTTP 401/403 behavior, logout/device ownership, selected audit and QR endpoints, and direct server-side validation bypass. Use xUnit service tests, server-log inspection and Swagger/backend-source comparison as supporting evidence. Controlled fixtures and evidence are recorded for every executed case.")
set_cell(security.cell(3, 1), "All 19 Security_CrossCutting cases have execution status and evidence: 15 Pass, 4 Fail, 0 Blocked and 0 Not Run. Confirmed authorization defects are recorded by BUG-ID; completion does not imply zero Critical/High defects.")

# Level 2 Integration Test table.
integration = doc.tables[15]
set_cell(integration.cell(5, 1), "SU26_StockLite_Report_5.2_IntegrationTests_L2.xlsm")
set_cell(integration.cell(7, 1), "Current result: 905 total L2 cases; 843 Pass, 62 Fail, 0 Blocked and 0 Not Run (93.1% overall pass rate). Controller/API scope: 849 cases; 795 Pass, 54 Fail, 0 Blocked and 0 Not Run (93.6% pass rate). Scheduler: 37 cases; 33 Pass and 4 Fail. Security: 19 cases; 15 Pass and 4 Fail. Fail results require BUG-ID evidence.")

# L2 environment row: remove obsolete unavailable-fixture/Blocked statement.
environment = doc.tables[21]
set_cell(environment.cell(2, 3), "Disposable Admin, Warehouse, Sales and Auditor sessions; credentials supplied through local environment variables and excluded from evidence. Required role, concurrency and fault-injection scenarios were completed with controlled local fixtures and evidence.")

# Target-account status: named accounts were not provisioned, but L2 execution was closed using controlled sessions.
accounts = doc.tables[22]
for row in (3, 4, 5):
    set_cell(accounts.cell(row, 3), "Covered by controlled L2 session")

# Current Report 5.2 artifact is macro-enabled.
background_level = doc.tables[16]
set_cell(background_level.cell(5, 1), "Included in SU26_StockLite_Report_5.2_IntegrationTests_L2.xlsm on dedicated job sheets.")
workbooks = doc.tables[25]
set_cell(workbooks.cell(2, 1), "SU26_StockLite_Report_5.2_IntegrationTests_L2.xlsm")

# Report summary: update only the L2 Integration column.
summary = doc.tables[27]
set_cell(summary.cell(2, 3), "843")
set_cell(summary.cell(3, 3), "62 Fail / 0 Blocked")
set_cell(summary.cell(4, 3), "93.1% overall (93.6% Controller/API)")

# Defect and work-product summaries: keep defect classification, update L2 execution denominator/status.
defects = doc.tables[28]
set_cell(defects.cell(1, 3), "48 distinct BUG IDs identified from 62 failed execution cases")

products = doc.tables[29]
set_cell(products.cell(1, 3), "48 distinct L2 BUG IDs from 62 failed cases; 0 cases remain Blocked.")

doc.save(out)
print(out)
